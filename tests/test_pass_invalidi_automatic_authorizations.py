import os
import shutil
import subprocess
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from core.file_state import capture_file_snapshot
from core.pass_invalidi_authorizations import (
    authorization_payload_for_record, authorization_replacements,
    find_authorization_doc_for_record, generate_authorization_doc_with_word,
)
from core.pass_invalidi_details import enrich_records, load_details, save_details


def record(**changes):
    value = {"numero": 61, "source": "REGISTRO INVALIDI COMUNE 2026.xlsx", "nome": "DE ROSSI MARIA",
             "cognome": "DE ROSSI", "nome_proprio": "Maria", "nato_il": "31/05/1980", "sesso": "F",
             "luogo_nascita": "Suzzara (MN)", "comune_residenza": "Pegognaga", "indirizzo": "via Roma 1",
             "rilascio": "13/10/2026", "scadenza": "31/05/2032", "protocol_number": "14064",
             "protocol_date": "02/10/2026", "issue_date": "13/10/2026", "note": "",
             "mode": "append", "_create_authorization": True}
    return dict(value, **changes)


def test_protocol_release_expiry_and_creation_remain_distinct():
    payload = authorization_payload_for_record(record(protocol_date="31/12/2025"))
    assert payload["year"] == 2026
    assert payload["release_date"] == "13/10/2026"
    assert payload["protocol_date"] == "31/12/2025"
    assert payload["expiry_date"] == "31/05/2032"
    replacements = authorization_replacements(payload)
    assert replacements["DESTINATARIO"] == "Alla Sig.ra DE ROSSI Maria"
    assert replacements["NATO_NATA"] == "nata"
    assert replacements["DATA_DOCUMENTO"] == "13/10/2026"
    assert "Validità" in replacements["RILASCIO_VALIDITA"]
    male = authorization_replacements(authorization_payload_for_record(record(sesso="M", nome_proprio="Mario")))
    assert male["DESTINATARIO"] == "Al Sig. DE ROSSI Mario" and male["NATO_NATA"] == "nato"


@pytest.mark.parametrize("changes", [{"sesso": ""}, {"protocol_number": ""}, {"nato_il": "non valida"},
                                     {"luogo_nascita": ""}, {"scadenza": "01/01/2020"}])
def test_missing_or_inconsistent_fields_do_not_create_authorization(changes):
    with pytest.raises(ValueError):
        authorization_payload_for_record(record(**changes))


def test_additional_fields_survive_reload_and_other_register_updates(tmp_path):
    path = tmp_path / "details.json"
    first = record()
    second = record(numero=61, source="registro 2025.xlsx", nome="BIANCHI MARIO", cognome="BIANCHI", nome_proprio="Mario")
    save_details([first], path)
    save_details([second], path)
    fresh = [{"numero": 61, "source": first["source"], "nome": first["nome"], "rilascio": "14/10/2026"}]
    enrich_records(fresh, path)
    assert fresh[0]["nato_il"] == "31/05/1980"
    assert fresh[0]["cognome"] == "DE ROSSI"
    assert fresh[0]["rilascio"] == "14/10/2026"  # Excel remains authoritative.
    assert len(load_details(path)["records"]) == 2
    reused = [{"numero": 61, "source": first["source"], "nome": "NUOVO TITOLARE"}]
    enrich_records(reused, path)
    assert "nato_il" not in reused[0]
    save_details([dict(first, mode="delete_existing")], path)
    assert len(load_details(path)["records"]) == 1


def test_corrupt_details_are_never_overwritten(tmp_path):
    path = tmp_path / "details.json"
    path.write_text("corrotto")
    with pytest.raises(RuntimeError):
        save_details([record()], path)
    assert path.read_text() == "corrotto"


def test_lookup_never_opens_a_different_person_or_partial_number(tmp_path):
    (tmp_path / "DE ROSSI Maria Autoriz. disabili n. 61 2026.doc").write_bytes(b"doc")
    assert find_authorization_doc_for_record(record(numero=1), tmp_path) is None
    assert find_authorization_doc_for_record(record(nome="VERDI ANNA", cognome="VERDI", nome_proprio="Anna"), tmp_path) is None


def test_word_failure_preserves_existing_file_and_cleans_temporary_output(tmp_path):
    template = tmp_path / "model.doc"
    target = tmp_path / "authorization.doc"
    template.write_bytes(b"template")
    target.write_bytes(b"previous authorization")
    def fail(args, **_kwargs):
        Path(args[args.index("-OutputPath") + 1]).write_bytes(b"incomplete")
        return subprocess.CompletedProcess(args, 1, "", "Word error")
    with patch("core.pass_invalidi_authorizations.subprocess.run", side_effect=fail), pytest.raises(RuntimeError):
        generate_authorization_doc_with_word(template, target, authorization_payload_for_record(record()), overwrite=True)
    assert target.read_bytes() == b"previous authorization"
    assert not list(tmp_path.glob("pass_auth_*"))
    with patch("core.pass_invalidi_authorizations.subprocess.run") as run, pytest.raises(FileExistsError):
        generate_authorization_doc_with_word(template, target, authorization_payload_for_record(record()))
    run.assert_not_called()


@pytest.fixture
def page(tmp_path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    from qt_app.pass_invalidi import PassInvalidiPage
    app = QApplication.instance() or QApplication([])
    original = tmp_path / record()["source"]
    wb = openpyxl.Workbook()
    wb.active.append(["N", "GENERALITA", "RILASCIO", "SCADENZA", "NOTE"])
    wb.save(original)
    working = tmp_path / "work.xlsx"
    shutil.copy2(original, working)
    config = {"paths": {"pass_invalidi_network_folder": str(tmp_path)}}
    widget = PassInvalidiPage(config)
    widget.primary_source_file = str(original)
    widget.working_copy_file = working
    widget.source_file_snapshot = capture_file_snapshot(original)
    widget.records = [record(_pending=True)]
    widget.pending_records = [record()]
    with patch("qt_app.pass_invalidi.QMessageBox.information"), patch("qt_app.pass_invalidi.QMessageBox.critical"), \
         patch("qt_app.pass_invalidi.log_audit_event"), patch("qt_app.pass_invalidi.create_excel_backup"):
        yield widget, original, app


def test_word_is_requested_only_after_excel_and_additional_fields_are_saved(page):
    widget, original, _app = page
    def check_saved(requests):
        wb = openpyxl.load_workbook(original)
        assert wb.active.cell(2, 1).value == 61
        assert wb.active.cell(2, 2).value == "DE ROSSI MARIA - VIA ROMA 1"
        wb.close()
        details = load_details(original.parent / "dati_autorizzazioni_pass.json")
        assert next(iter(details["records"].values()))["protocol_number"] == "14064"
        assert requests[0]["numero"] == 61
    with patch.object(widget, "_start_authorizations", side_effect=check_saved) as generate:
        assert widget.save_changes(trigger_reload=False)
    generate.assert_called_once()
    assert widget.pending_records == []


def test_excel_failure_never_requests_word(page):
    widget, original, _app = page
    before = original.read_bytes()
    with patch("qt_app.pass_invalidi.write_pending_to_xlsx", side_effect=OSError("File locked")), patch.object(widget, "_start_authorizations") as generate:
        assert not widget.save_changes(trigger_reload=False)
    assert original.read_bytes() == before
    assert widget.pending_records
    generate.assert_not_called()


def test_details_failure_can_be_retried_after_excel_was_saved(page):
    widget, _original, _app = page
    with patch("qt_app.pass_invalidi.save_details", side_effect=OSError("Network unavailable")), patch.object(widget, "_start_authorizations") as generate:
        assert not widget.save_changes(trigger_reload=False)
    generate.assert_not_called()
    with patch.object(widget, "_start_authorizations") as generate:
        assert widget.save_changes(trigger_reload=False)
    generate.assert_called_once()


def test_creation_dialog_preserves_register_expiry_and_requires_explicit_sex():
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    from qt_app.pass_invalidi import PassInvalidiAuthorizationDialog
    app = QApplication.instance() or QApplication([])
    dialog = PassInvalidiAuthorizationDialog(record())
    assert dialog.data_scadenza.text() == "31/05/2032"
    dialog.protocollo_data.setText("31/12/2025")
    assert dialog.data_scadenza.text() == "31/05/2032"
    assert dialog.cognome.text() == "DE ROSSI"
    assert dialog.data_rilascio.isReadOnly() and dialog.data_scadenza.isReadOnly()
    with patch("qt_app.pass_invalidi.QMessageBox.warning"):
        dialog.sesso.setCurrentText("")
        dialog._accept()
    assert dialog.payload_input is None
    assert app
